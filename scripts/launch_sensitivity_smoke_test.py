import os
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest


ROOT = Path(__file__).resolve().parents[1]
APP_PATH = ROOT / "streamlit_app.py"
MODEL_PATH = ROOT / "demo-data" / "Launch_Master_Model.xlsx"
os.chdir(ROOT)
sys.path.insert(0, str(ROOT))


class TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.row = None
        self.cell = None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag in {"th", "td"} and self.row is not None:
            self.cell = []

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)

    def handle_endtag(self, tag):
        if tag in {"th", "td"} and self.cell is not None:
            self.row.append("".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def assert_clean(app, label):
    if list(app.exception):
        raise AssertionError(f"{label}: {app.exception[0].message}")


def widget_by_key(widgets, key):
    return next(widget for widget in widgets if widget.key == key)


def scenario_rows(app):
    html = next(
        str(item.value)
        for item in app.markdown
        if "<table" in str(item.value) and "5Y NPV" in str(item.value) and "Payback Period" in str(item.value)
    )
    parser = TableParser()
    parser.feed(html)
    headers = parser.rows[0]
    return {row[0]: dict(zip(headers, row)) for row in parser.rows[1:]}


def number(value):
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    return float(cleaned or 0)


def open_sensitivity(case_id="LAUNCH-1002"):
    app = AppTest.from_file(str(APP_PATH), default_timeout=120)
    app.session_state["current_module"] = "launch"
    app.session_state["launch_page"] = "Launch Case"
    app.session_state["selected_launch_case_id"] = case_id
    app.session_state["launch_current_role"] = "Marketing · Launch Coordinator"
    app.session_state[f"launch_case_section_{case_id}"] = "Sensitivity"
    app.run()
    assert_clean(app, f"{case_id} Sensitivity")
    return app


workbook = load_workbook(MODEL_PATH, read_only=True, data_only=False)
assert workbook.sheetnames == ["Reference Data", "Base Inputs", "Sensitivity Inputs", "Calculation Engine", "P&L", "Outputs"]
assert workbook["Outputs"]["B6"].value == "=NPV('Reference Data'!B2,'P&L'!H2:H6)"
assert str(workbook["Outputs"]["B7"].value).startswith("=IF('P&L'!I2>=0")
workbook.close()

app = open_sensitivity()
page_text = "\n".join(str(item.value) for item in app.markdown)
for label in (
    "Market Share",
    "Treatment Eligibility",
    "Sales Coverage",
    "Access Reach",
    "Net Price",
    "Regulatory Timing",
    "COGS",
    "Discount Rate",
):
    assert label in page_text, f"Missing Sensitivity content: {label}"
methodology_labels = {expander.label for expander in app.expander}
for driver in ("Market Share", "Treatment Eligibility", "Sales Coverage", "Patient Access Reach", "Net Price", "Regulatory Timing", "COGS", "Discount Rate"):
    assert f"{driver} methodology" in methodology_labels
assert "Y5 Target sets the end-state sensitivity; intermediate years are ramped automatically." in {
    str(item.value) for item in app.caption
}
assert any("Application" in str(item.value) for item in app.markdown if "<table" in str(item.value))
treatment_base = next(widget for widget in app.text_input if widget.label == "Treatment Eligibility Base")
assert treatment_base.value == "68.0%", treatment_base.value

initial = scenario_rows(app)
base_before = dict(initial["Base"])
market_share = widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_market_share_downside_terminal")
market_share.set_value(-10.0).run()
assert_clean(app, "Market Share downside change")
market_changed = scenario_rows(app)
assert market_changed["Base"] == base_before
for metric in ("Y5 Patients on Product", "Y5 Net Revenue", "Y5 Operating Profit", "5Y NPV"):
    assert number(market_changed["Downside"][metric]) < number(initial["Downside"][metric]), metric

for driver_key, new_value in (
    ("treatment_eligibility", -20.0),
    ("sales_coverage", -20.0),
    ("access_reach", -20.0),
):
    app = open_sensitivity()
    initial = scenario_rows(app)
    widget_by_key(app.number_input, f"launch_sensitivity_LAUNCH-1002_{driver_key}_downside_terminal").set_value(new_value).run()
    assert_clean(app, f"{driver_key} downside change")
    changed = scenario_rows(app)
    assert changed["Base"] == initial["Base"]
    for metric in ("Y5 Patients on Product", "Y5 Net Revenue", "Y5 Operating Profit", "5Y NPV"):
        assert number(changed["Downside"][metric]) < number(initial["Downside"][metric]), (driver_key, metric)

app = open_sensitivity()
initial = scenario_rows(app)
widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_net_price_downside_terminal").set_value(-20.0).run()
assert_clean(app, "Net Price downside change")
price_changed = scenario_rows(app)
assert price_changed["Base"] == initial["Base"]
assert price_changed["Downside"]["Y5 Patients on Product"] == initial["Downside"]["Y5 Patients on Product"]
for metric in ("Y5 Net Revenue", "Y5 Operating Profit", "5Y NPV"):
    assert number(price_changed["Downside"][metric]) < number(initial["Downside"][metric]), metric

app = open_sensitivity()
initial = scenario_rows(app)
widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_regulatory_timing_downside_terminal").set_value(365.0).run()
assert_clean(app, "Regulatory Timing downside change")
timing_changed = scenario_rows(app)
assert timing_changed["Base"] == initial["Base"]
assert number(timing_changed["Downside"]["5Y Cumulative Revenue"]) < number(initial["Downside"]["5Y Cumulative Revenue"])

app = open_sensitivity()
initial = scenario_rows(app)
cogs = widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_cogs_downside_terminal")
cogs.set_value(20.0).run()
assert_clean(app, "COGS downside change")
cogs_changed = scenario_rows(app)
assert cogs_changed["Base"] == initial["Base"]
for metric in ("Y5 Patients on Product", "Y5 Net Revenue", "5Y Cumulative Revenue"):
    assert cogs_changed["Downside"][metric] == initial["Downside"][metric], metric
for metric in ("Y5 Gross Margin %", "Y5 Operating Profit", "5Y NPV"):
    assert number(cogs_changed["Downside"][metric]) < number(initial["Downside"][metric]), metric

app = open_sensitivity()
initial = scenario_rows(app)
discount = widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_discount_rate_downside_terminal")
discount.set_value(4.0).run()
assert_clean(app, "Discount Rate downside change")
discount_changed = scenario_rows(app)
assert discount_changed["Base"] == initial["Base"]
for metric in (
    "Y5 Patients on Product",
    "Y5 Net Revenue",
    "5Y Cumulative Revenue",
    "Y5 Gross Margin %",
    "Y5 Operating Profit",
    "5Y Cumulative Operating Profit",
    "Y5 Operating Margin %",
    "Payback Period",
):
    assert discount_changed["Downside"][metric] == initial["Downside"][metric], metric
assert number(discount_changed["Downside"]["5Y NPV"]) != number(initial["Downside"]["5Y NPV"]), (
    initial["Downside"]["5Y NPV"],
    discount_changed["Downside"]["5Y NPV"],
)

mode = widget_by_key(app.radio, "launch_sensitivity_mode_LAUNCH-1002")
mode.set_value("By Year").run()
assert_clean(app, "By Year mode")
for year in ("Y1", "Y2", "Y3", "Y4", "Y5"):
    widget_by_key(app.number_input, f"launch_sensitivity_LAUNCH-1002_market_share_downside_{year.lower()}")
by_year_before = scenario_rows(app)
widget_by_key(app.number_input, "launch_sensitivity_LAUNCH-1002_market_share_downside_y5").set_value(-15.0).run()
assert_clean(app, "By Year Market Share change")
by_year_after = scenario_rows(app)
assert by_year_after["Base"] == by_year_before["Base"]
assert number(by_year_after["Downside"]["Y5 Patients on Product"]) < number(by_year_before["Downside"]["Y5 Patients on Product"])

# One edit in By Year mode must persist and feed recalculation for every year-based driver.
for driver_key, new_value in (
    ("treatment_eligibility", -25.0),
    ("sales_coverage", -25.0),
    ("access_reach", -25.0),
    ("net_price", -25.0),
    ("cogs", 25.0),
):
    app = open_sensitivity()
    widget_by_key(app.radio, "launch_sensitivity_mode_LAUNCH-1002").set_value("By Year").run()
    assert_clean(app, f"{driver_key} By Year mode")
    before = scenario_rows(app)
    key = f"launch_sensitivity_LAUNCH-1002_{driver_key}_downside_y5"
    widget_by_key(app.number_input, key).set_value(new_value).run()
    assert_clean(app, f"{driver_key} By Year first edit")
    after = scenario_rows(app)
    assert after["Base"] == before["Base"]
    assert widget_by_key(app.number_input, key).value == new_value
    assert after["Downside"] != before["Downside"], driver_key

print("launch sensitivity smoke test complete")
