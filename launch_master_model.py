from __future__ import annotations

from copy import deepcopy
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


LAUNCH_YEARS = ("Y1", "Y2", "Y3", "Y4", "Y5")
SCENARIOS = ("Downside", "Base", "Upside")


@dataclass(frozen=True)
class SensitivityDriver:
    name: str
    owner: str
    unit: str
    method: str
    model_input: str
    affects: str
    downstream: str
    default_downside: float
    default_upside: float


@lru_cache(maxsize=4)
def load_sensitivity_drivers(workbook_path: str) -> tuple[SensitivityDriver, ...]:
    """Read controlled sensitivity definitions from the master workbook."""
    workbook = load_workbook(Path(workbook_path), read_only=True, data_only=False)
    sheet = workbook["Sensitivity Inputs"]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    positions = {name: index for index, name in enumerate(headers)}
    required = {
        "Driver",
        "Owner",
        "Unit",
        "Method",
        "Model Input",
        "Affects",
        "Downstream",
        "Default Downside",
        "Default Upside",
    }
    missing = required - positions.keys()
    if missing:
        raise ValueError(f"Launch master model is missing sensitivity fields: {', '.join(sorted(missing))}")

    definitions = []
    for row in rows:
        name = str(row[positions["Driver"]] or "").strip()
        if not name:
            continue
        definitions.append(
            SensitivityDriver(
                name=name,
                owner=str(row[positions["Owner"]] or "").strip(),
                unit=str(row[positions["Unit"]] or "").strip(),
                method=str(row[positions["Method"]] or "").strip(),
                model_input=str(row[positions["Model Input"]] or "").strip(),
                affects=str(row[positions["Affects"]] or "").strip(),
                downstream=str(row[positions["Downstream"]] or "").strip(),
                default_downside=float(row[positions["Default Downside"]] or 0),
                default_upside=float(row[positions["Default Upside"]] or 0),
            )
        )
    workbook.close()
    if not definitions:
        raise ValueError("Launch master model contains no sensitivity driver definitions.")
    return tuple(definitions)


def _ramp_weights(first_commercial_year: str) -> dict[str, float]:
    try:
        start = LAUNCH_YEARS.index(first_commercial_year)
    except ValueError:
        start = 0
    periods = len(LAUNCH_YEARS) - start
    return {
        year: 0.0 if index < start else (index - start + 1) / periods
        for index, year in enumerate(LAUNCH_YEARS)
    }


def default_sensitivity_settings(drivers: tuple[SensitivityDriver, ...]) -> dict[str, object]:
    settings: dict[str, object] = {"Mode": "Terminal / Y5", "Drivers": {}, "Completed Drivers": []}
    driver_settings = settings["Drivers"]
    assert isinstance(driver_settings, dict)
    for driver in drivers:
        weights = _ramp_weights("Y1") if driver.method in {"RAMP_PP", "RAMP_RELATIVE"} else {year: 1.0 for year in LAUNCH_YEARS}
        driver_settings[driver.name] = {
            "Downside": driver.default_downside,
            "Upside": driver.default_upside,
            "By Year": {
                scenario: {year: value * weights[year] for year in LAUNCH_YEARS}
                for scenario, value in {
                    "Downside": driver.default_downside,
                    "Upside": driver.default_upside,
                }.items()
            },
        }
    return settings


def normalize_sensitivity_settings(
    current: object,
    drivers: tuple[SensitivityDriver, ...],
) -> dict[str, object]:
    defaults = default_sensitivity_settings(drivers)
    if not isinstance(current, dict):
        return defaults
    mode = str(current.get("Mode", "Terminal / Y5"))
    normalized = deepcopy(defaults)
    normalized["Mode"] = mode if mode in {"Terminal / Y5", "By Year"} else "Terminal / Y5"
    completed = current.get("Completed Drivers", [])
    if isinstance(completed, list):
        valid_names = {driver.name for driver in drivers}
        normalized["Completed Drivers"] = [str(name) for name in completed if str(name) in valid_names]
    source_drivers = current.get("Drivers", current)
    if not isinstance(source_drivers, dict):
        return normalized
    normalized_drivers = normalized["Drivers"]
    assert isinstance(normalized_drivers, dict)
    for driver in drivers:
        source = source_drivers.get(driver.name, {})
        if not isinstance(source, dict):
            continue
        target = normalized_drivers[driver.name]
        for scenario in ("Downside", "Upside"):
            try:
                target[scenario] = float(source.get(scenario, target[scenario]))
            except (TypeError, ValueError):
                pass
        source_by_year = source.get("By Year", {})
        if not isinstance(source_by_year, dict):
            continue
        for scenario in ("Downside", "Upside"):
            scenario_values = source_by_year.get(scenario, {})
            if not isinstance(scenario_values, dict):
                continue
            for year in LAUNCH_YEARS:
                try:
                    target["By Year"][scenario][year] = float(
                        scenario_values.get(year, target["By Year"][scenario][year])
                    )
                except (TypeError, ValueError):
                    pass
    return normalized


def _deviation_by_year(
    driver: SensitivityDriver,
    settings: dict[str, object],
    scenario: str,
    first_commercial_year: str,
) -> dict[str, float]:
    driver_settings = settings["Drivers"][driver.name]
    if settings.get("Mode") == "By Year" and driver.method not in {"DAYS_SHIFT", "NPV_PP"}:
        return {
            year: float(driver_settings["By Year"][scenario][year])
            for year in LAUNCH_YEARS
        }
    terminal = float(driver_settings[scenario])
    if driver.method in {"RAMP_PP", "RAMP_RELATIVE"}:
        weights = _ramp_weights(first_commercial_year)
        return {year: terminal * weights[year] for year in LAUNCH_YEARS}
    return {year: terminal for year in LAUNCH_YEARS}


def build_scenario_inputs(
    base_inputs: dict[str, object],
    settings: dict[str, object],
    drivers: tuple[SensitivityDriver, ...],
    scenario: str,
    first_commercial_year: str,
) -> dict[str, object]:
    """Translate workbook-defined deviations into model inputs and an audit trace."""
    if scenario == "Base":
        return {
            "adjustments": {},
            "discount_rate": float(base_inputs.get("Discount Rate", 0)),
            "trace": [],
            "warnings": [],
        }

    adjustments: dict[str, object] = {}
    trace: list[dict[str, object]] = []
    warnings: list[str] = []
    discount_rate = float(base_inputs.get("Discount Rate", 0))
    for driver in drivers:
        deviations = _deviation_by_year(driver, settings, scenario, first_commercial_year)
        base = base_inputs.get(driver.name, {})
        base_by_year = base if isinstance(base, dict) else {year: base for year in LAUNCH_YEARS}
        scenario_values: dict[str, float] = {}

        if driver.method == "NPV_PP":
            discount_rate = max(0.0, discount_rate + deviations["Y5"] / 100.0)
            scenario_values = {year: discount_rate for year in LAUNCH_YEARS}
        elif driver.method == "DAYS_SHIFT":
            adjustments[driver.model_input] = deviations
            scenario_values = {
                year: float(base_by_year.get(year, 0) or 0) + deviations[year]
                for year in LAUNCH_YEARS
            }
        elif driver.method in {"ALL_YEARS_RELATIVE", "RAMP_RELATIVE"}:
            factors = {year: max(0.0, 1.0 + deviations[year] / 100.0) for year in LAUNCH_YEARS}
            if driver.name == "Sales Coverage" and not any(float(base_by_year.get(year, 0) or 0) > 0 for year in LAUNCH_YEARS):
                factors = {year: 1.0 for year in LAUNCH_YEARS}
                warnings.append("Sales Coverage has no active Base value, so no adoption impact was calculated.")
            adjustments[driver.model_input] = factors
            scenario_values = {
                year: min(1.0, max(0.0, float(base_by_year.get(year, 0) or 0) * factors[year]))
                if driver.name in {"Treatment Eligibility", "Sales Coverage"}
                else float(base_by_year.get(year, 0) or 0) * factors[year]
                for year in LAUNCH_YEARS
            }
        elif driver.method == "RAMP_PP":
            deltas = {year: deviations[year] / 100.0 for year in LAUNCH_YEARS}
            adjustments[driver.model_input] = deltas
            scenario_values = {
                year: min(1.0, max(0.0, float(base_by_year.get(year, 0) or 0) + deltas[year]))
                for year in LAUNCH_YEARS
            }
        else:
            raise ValueError(f"Unsupported master-model sensitivity method: {driver.method}")

        for year in LAUNCH_YEARS:
            trace.append(
                {
                    "Scenario": scenario,
                    "Driver": driver.name,
                    "Year": year,
                    "Base Workstream Input": base_by_year.get(year, base),
                    "Sensitivity Adjustment": deviations[year],
                    "Scenario Input": scenario_values[year],
                    "Model Input": driver.model_input,
                }
            )
    return {
        "adjustments": adjustments,
        "discount_rate": discount_rate,
        "trace": trace,
        "warnings": warnings,
    }


def calculate_management_valuation(
    annual_operating_profit: dict[str, float],
    annual_non_operating_outflows: dict[str, float],
    discount_rate: float,
) -> tuple[float, str, dict[str, float]]:
    """Demo adapter for the workbook-defined management NPV and payback formulas."""
    cash_flows = {
        year: float(annual_operating_profit.get(year, 0))
        - float(annual_non_operating_outflows.get(year, 0))
        for year in LAUNCH_YEARS
    }
    rate = max(0.0, float(discount_rate))
    npv = sum(cash_flows[year] / ((1.0 + rate) ** index) for index, year in enumerate(LAUNCH_YEARS, start=1))
    cumulative = 0.0
    payback = ">5 years"
    for index, year in enumerate(LAUNCH_YEARS, start=1):
        prior = cumulative
        current = cash_flows[year]
        cumulative += current
        if cumulative >= 0 and current > 0:
            crossing = (index - 1) + max(0.0, min(1.0, -prior / current))
            payback = "<1 year" if crossing < 1 else f"{crossing:.1f} years"
            break
    return npv, payback, cash_flows


def evaluate_scenarios(
    base_inputs: dict[str, object],
    settings: dict[str, object],
    drivers: tuple[SensitivityDriver, ...],
    first_commercial_year: str,
    calculate_model: Callable[[dict[str, object]], dict[str, object]],
) -> dict[str, dict[str, object]]:
    """Run the same integrated model once per analytical overlay."""
    results = {}
    for scenario in SCENARIOS:
        scenario_inputs = build_scenario_inputs(
            base_inputs,
            settings,
            drivers,
            scenario,
            first_commercial_year,
        )
        model = calculate_model(scenario_inputs["adjustments"])
        results[scenario] = {**scenario_inputs, "model": model}
    return results
